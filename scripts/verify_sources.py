"""Verify source files required for Day 2 data processing.

Locates MetroPT-3, AI4I 2020, UCI Online Retail, and M5 (optional).
Prints resolved paths, file sizes, and row/column counts.
Exits non-zero if any required source file is missing.
"""

from __future__ import annotations

import argparse
import logging
import sys
from pathlib import Path

import openpyxl
import pandas as pd

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger(__name__)

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "raw"
DOCS_DIR = ROOT / "docs"


def count_csv_rows_and_cols(
    file_path: Path, chunksize: int = 100_000
) -> tuple[int, int, list[str]]:
    """Count rows and columns of a CSV using chunked reading to conserve memory."""
    total_rows = 0
    cols: list[str] = []
    for chunk in pd.read_csv(file_path, chunksize=chunksize, low_memory=False):
        if not cols:
            cols = list(chunk.columns)
        total_rows += len(chunk)
    return total_rows, len(cols), cols


def verify_sources(output_file: Path | None = None) -> int:
    report_lines: list[str] = []

    def log_and_record(msg: str) -> None:
        print(msg)
        report_lines.append(msg)

    log_and_record("=" * 70)
    log_and_record("AI Operations Intelligence Platform - Source Verification Report")
    log_and_record("=" * 70)

    missing_required: list[str] = []

    # 1. MetroPT-3
    log_and_record("\n[1] MetroPT-3 (Air Compressor)")
    metro_dir = RAW_DIR / "metropt"
    metro_csv_matches = list(metro_dir.glob("*.csv"))
    metro_pdf_matches = list(metro_dir.glob("*.pdf"))

    if not metro_csv_matches:
        log_and_record("  FAIL: MetroPT CSV not found in data/raw/metropt/")
        missing_required.append("MetroPT CSV")
    else:
        metro_csv = metro_csv_matches[0]
        size_mb = metro_csv.stat().st_size / (1024 * 1024)
        log_and_record(f"  File: {metro_csv.name}")
        log_and_record(f"  Path: {metro_csv.resolve()}")
        log_and_record(f"  Size: {size_mb:.2f} MB ({metro_csv.stat().st_size:,} bytes)")
        try:
            rows, cols, col_names = count_csv_rows_and_cols(metro_csv)
            log_and_record(f"  Rows: {rows:,}, Columns: {cols}")
            log_and_record(f"  Columns: {col_names}")
        except Exception as e:
            log_and_record(f"  Error reading CSV: {e}")
            missing_required.append("MetroPT CSV (read error)")

    if not metro_pdf_matches:
        log_and_record("  FAIL: Data Description_Metro.pdf not found in data/raw/metropt/")
        missing_required.append("Data Description_Metro.pdf")
    else:
        metro_pdf = metro_pdf_matches[0]
        size_kb = metro_pdf.stat().st_size / 1024
        log_and_record(f"  PDF : {metro_pdf.name}")
        log_and_record(f"  Path: {metro_pdf.resolve()}")
        log_and_record(f"  Size: {size_kb:.2f} KB ({metro_pdf.stat().st_size:,} bytes)")

    # 2. AI4I 2020
    log_and_record("\n[2] AI4I 2020 Predictive Maintenance")
    ai4i_dir = RAW_DIR / "ai4i"
    ai4i_matches = list(ai4i_dir.glob("*.csv"))
    if not ai4i_matches:
        log_and_record("  FAIL: ai4i2020.csv not found in data/raw/ai4i/")
        missing_required.append("ai4i2020.csv")
    else:
        ai4i_csv = ai4i_matches[0]
        size_kb = ai4i_csv.stat().st_size / 1024
        log_and_record(f"  File: {ai4i_csv.name}")
        log_and_record(f"  Path: {ai4i_csv.resolve()}")
        log_and_record(f"  Size: {size_kb:.2f} KB ({ai4i_csv.stat().st_size:,} bytes)")
        try:
            df_ai4i = pd.read_csv(ai4i_csv)
            log_and_record(f"  Rows: {len(df_ai4i):,}, Columns: {df_ai4i.shape[1]}")
            log_and_record(f"  Columns: {list(df_ai4i.columns)}")
        except Exception as e:
            log_and_record(f"  Error reading CSV: {e}")
            missing_required.append("ai4i2020.csv (read error)")

    # 3. UCI Online Retail
    log_and_record("\n[3] UCI Online Retail")
    retail_dir = RAW_DIR / "retail"
    retail_matches = list(retail_dir.glob("*.xlsx")) + list(retail_dir.glob("*.csv"))
    if not retail_matches:
        log_and_record("  FAIL: Online Retail.xlsx not found in data/raw/retail/")
        missing_required.append("Online Retail.xlsx")
    else:
        retail_file = retail_matches[0]
        size_mb = retail_file.stat().st_size / (1024 * 1024)
        log_and_record(f"  File: {retail_file.name}")
        log_and_record(f"  Path: {retail_file.resolve()}")
        log_and_record(f"  Size: {size_mb:.2f} MB ({retail_file.stat().st_size:,} bytes)")
        try:
            if retail_file.suffix.lower() == ".xlsx":
                wb = openpyxl.load_workbook(retail_file, read_only=True)
                sheet_names = wb.sheetnames
                active_sheet = wb[sheet_names[0]]
                max_r = active_sheet.max_row or "Unknown"
                max_c = active_sheet.max_column or "Unknown"
                wb.close()
                log_and_record(f"  Sheets: {sheet_names}")
                log_and_record(
                    f"  Active Sheet: {sheet_names[0]} (Approx Rows: {max_r}, Cols: {max_c})"
                )
            else:
                rows, cols, col_names = count_csv_rows_and_cols(retail_file)
                log_and_record(f"  Rows: {rows:,}, Columns: {cols}")
                log_and_record(f"  Columns: {col_names}")
        except Exception as e:
            log_and_record(f"  Error reading file: {e}")
            missing_required.append(f"{retail_file.name} (read error)")

    # 4. M5 Forecasting (Optional)
    log_and_record("\n[4] M5 Forecasting (Optional)")
    m5_dir = RAW_DIR / "m5"
    m5_files = ["calendar.csv", "sales_train_validation.csv", "sell_prices.csv"]
    m5_present = [f for f in m5_files if (m5_dir / f).exists()]
    if len(m5_present) == len(m5_files):
        log_and_record("  Status: All M5 files present.")
        for f in m5_files:
            fp = m5_dir / f
            size_mb = fp.stat().st_size / (1024 * 1024)
            log_and_record(f"    - {f}: {size_mb:.2f} MB ({fp.resolve()})")
    elif m5_present:
        log_and_record(f"  Status: Partial M5 files present: {m5_present}")
    else:
        log_and_record(
            "  Status: M5 files not found (optional - task T9 will be skipped with warning)."
        )

    log_and_record("\n" + "=" * 70)
    if missing_required:
        log_and_record(
            f"VERIFICATION FAILED: Missing required source(s): {', '.join(missing_required)}"
        )
        log_and_record("=" * 70)
    else:
        log_and_record("VERIFICATION PASSED: All required source files are present and readable.")
        log_and_record("=" * 70)

    # Save report
    if output_file:
        output_file.parent.mkdir(parents=True, exist_ok=True)
        output_file.write_text("\n".join(report_lines) + "\n", encoding="utf-8")
        logger.info("Report saved to %s", output_file)

    return 1 if missing_required else 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Verify raw data sources for Day 2.")
    parser.add_argument(
        "--output",
        type=Path,
        default=DOCS_DIR / "source_verification.txt",
        help="Path to save verification report (default: docs/source_verification.txt)",
    )
    args = parser.parse_args()
    sys.exit(verify_sources(output_file=args.output))
